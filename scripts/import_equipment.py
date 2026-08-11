"""Актуализация перечня оборудования из Excel «Актуальный перечень оборудования».

Excel — источник правды по парку: клещи (пистолеты, ID = G.NNN), трансформаторы
(ID = станция+-T#) и их станции/линии. Скрипт приводит справочники и привязки
БД в соответствие с файлом:

  ДОБАВЛЯЕТ  недостающие станции, трансформаторы (со связью со станцией), клещи;
  ИСПРАВЛЯЕТ тип клещей, где он в Excel другой (Excel вернее);
  ПЕРЕНОСИТ  клещ на трансформатор из Excel (старая привязка закрывается датой,
             создаётся новая active — как «перенос гана» в редакторе). Станция и
             линия пересчитываются автоматически по трансформатору.

НЕ УДАЛЯЕТ оборудование, которого нет в Excel (клещи/дубли/лишние трансформаторы,
станции) — только перечисляет в отчёте для ручного решения (на клещи могут
ссылаться уставки/дефекты).

Запуск:
    python scripts/import_equipment.py "путь/к/Актуальный перечень оборудовани.xlsx"           # dry-run
    python scripts/import_equipment.py "путь/к/…xlsx" --apply                                   # применить (бэкап!)

Аудит-триггеры на время bulk-загрузки снимаются (журнал не засоряется);
откат — из бэкапа БД. Приложение пересоздаст триггеры при старте.
"""
import re
import sqlite3
import sys
from datetime import date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from app.audit import drop_audit_triggers  # noqa: E402

DB_PATH = ROOT / 'data' / 'welding_shop.db'
TODAY = date.today().isoformat()
_COMMENT = 'актуализация из Excel'


def _norm(s) -> str:
    return re.sub(r'\s+', ' ', str(s or '').strip())


def _gnum(s):
    m = re.search(r'(\d+)', str(s or ''))
    return int(m.group(1)) if m else None


def parse_excel(path: str):
    """→ (guns, trans): guns[g_num]=(type, trans_code); trans[trans_code]=(AC|DC, station_code, brand)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    guns, trans = {}, {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(v is not None for v in r):
            continue
        name = _norm(r[1]).lower()
        if 'пистолет' in name:            # клещ: ID=G.NNN, ТИП, СТАНЦИЯ=трансформатор
            g = _gnum(r[0])
            if g is not None:
                guns[g] = (_norm(r[2]), _norm(r[6]))
        elif 'трансформатор' in name:     # трансформатор: ID, AC/DC, станция, бренд
            ac_dc = 'AC' if name.startswith('ac') else 'DC'
            trans[_norm(r[0])] = (ac_dc, _norm(r[6]), _norm(r[4]))
    wb.close()
    return guns, trans


def reconcile(db: sqlite3.Connection, guns_x: dict, trans_x: dict) -> dict:
    """Применить изменения к БД (в открытой транзакции). Вернуть отчёт (счётчики + списки)."""
    rep = {'station_add': 0, 'trans_add': 0, 'gun_add': 0, 'type_fix': [],
           'reassign': 0, 'unknown_trans': [], 'db_only_guns': [], 'dup_guns': [],
           'db_only_trans': [], 'db_only_stations': []}

    brand_id = {r[1]: r[0] for r in db.execute('SELECT UniqueID, brand FROM brand')}

    # станции по бренду — из трансформаторов Excel
    station_x = {st: br for (_ac, st, br) in trans_x.values() if st}
    station_id = {_norm(r[1]): r[0] for r in db.execute('SELECT UniqueID, station_name FROM station')}
    for st, br in station_x.items():
        if st not in station_id:
            cur = db.execute('INSERT INTO station (station_name, brand_id) VALUES (?,?)',
                             (st, brand_id.get(br)))
            station_id[st] = cur.lastrowid
            rep['station_add'] += 1

    # трансформаторы
    trans_id = {_norm(r[1]): r[0] for r in db.execute('SELECT UniqueID, transID FROM trans')}
    for tcode, (ac_dc, st, _br) in trans_x.items():
        if tcode not in trans_id:
            cur = db.execute('INSERT INTO trans (transID, type) VALUES (?,?)', (tcode, ac_dc))
            trans_id[tcode] = cur.lastrowid
            rep['trans_add'] += 1
            if st and st in station_id:  # привязка нового трансформатора к станции
                db.execute("INSERT INTO transformer_station_assignment "
                           "(start_date, end_date, is_active, comment, transformer_id, station_id) "
                           "VALUES (?, NULL, 1, ?, ?, ?)", (TODAY, _COMMENT, trans_id[tcode], station_id[st]))

    # клещи в БД (g_num -> список (id, type)); дубли не трогаем
    guns_db = {}
    for uid, g, ty in db.execute('SELECT UniqueID, g_num, gun_type FROM gun'):
        guns_db.setdefault(g, []).append((uid, ty))

    # активные привязки клещ->трансформатор
    gta = {r[0]: (r[1], r[2]) for r in db.execute(
        'SELECT gun_id, UniqueID, transformer_id FROM gun_transformer_assignment WHERE is_active=1')}

    for g in sorted(guns_x):
        gtype, tcode = guns_x[g]
        rows = guns_db.get(g)
        if rows is None:                                   # новый клещ
            cur = db.execute('INSERT INTO gun (g_num, gun_type) VALUES (?,?)', (g, gtype))
            gid = cur.lastrowid
        elif len(rows) > 1:                                # дубль — не трогаем
            rep['dup_guns'].append(g)
            continue
        else:
            gid, cur_type = rows[0]
            if gtype and cur_type != gtype:                # исправить тип
                db.execute('UPDATE gun SET gun_type=? WHERE UniqueID=?', (gtype, gid))
                rep['type_fix'].append((g, cur_type, gtype))

        # привязка к трансформатору из Excel
        if not tcode:
            continue
        if tcode not in trans_id:                          # клещ ссылается на неизвестный «трансформатор»
            rep['unknown_trans'].append((g, tcode))
            continue
        target_tid = trans_id[tcode]
        cur_assign = gta.get(gid)
        if cur_assign and cur_assign[1] == target_tid:
            continue                                       # уже на нужном
        if cur_assign:                                     # закрыть старую привязку датой
            db.execute('UPDATE gun_transformer_assignment SET is_active=0, end_date=? WHERE UniqueID=?',
                       (TODAY, cur_assign[0]))
        db.execute("INSERT INTO gun_transformer_assignment "
                   "(start_date, end_date, is_active, comments, gun_id, transformer_id) "
                   "VALUES (?, NULL, 1, ?, ?, ?)", (TODAY, _COMMENT, gid, target_tid))
        rep['reassign'] += 1

    # отчёт по тому, чего нет в Excel (не трогаем)
    rep['db_only_guns'] = sorted(set(guns_db) - set(guns_x))
    rep['dup_guns'] = sorted(set(rep['dup_guns']) | {g for g, v in guns_db.items() if len(v) > 1})
    rep['db_only_trans'] = sorted(set(trans_id) - set(trans_x))
    rep['db_only_stations'] = sorted(set(station_id) - set(station_x))
    return rep


def _print_report(rep: dict):
    print('\n=== ИЗМЕНЕНИЯ ===')
    print(f'  + станций:        {rep["station_add"]}')
    print(f'  + трансформаторов:{rep["trans_add"]}')
    print(f'  + клещей:         {rep["gun_add"]}')
    print(f'  ~ тип клещей:     {len(rep["type_fix"])}')
    for g, a, b in rep['type_fix']:
        print(f'      G.{g}: {a!r} → {b!r}')
    print(f'  ⇄ переносов клещ→трансформатор: {rep["reassign"]}')
    print('\n=== НЕ ТРОНУТО (нет в Excel) — решить вручную ===')
    print(f'  клещи только в БД: {len(rep["db_only_guns"])} → {rep["db_only_guns"]}')
    print(f'  дубли G-номера:    {len(rep["dup_guns"])} → {rep["dup_guns"]}')
    print(f'  трансформаторы только в БД: {len(rep["db_only_trans"])}')
    print(f'  станции только в БД:        {len(rep["db_only_stations"])}')
    if rep['unknown_trans']:
        print(f'  клещи со ссылкой на неизвестный трансформатор: {len(rep["unknown_trans"])} '
              f'→ {rep["unknown_trans"][:10]}')


def run(path: str, apply: bool):
    guns_x, trans_x = parse_excel(path)
    print(f'Excel: клещей {len(guns_x)}, трансформаторов {len(trans_x)}')
    db = sqlite3.connect(DB_PATH)
    try:
        drop_audit_triggers(db)
        # gun_add считаем отдельно: новые g_num
        existing = {r[0] for r in db.execute('SELECT DISTINCT g_num FROM gun')}
        rep = reconcile(db, guns_x, trans_x)
        rep['gun_add'] = len(set(guns_x) - existing)
        _print_report(rep)
        if apply:
            db.commit()
            print('\nПрименено. Перезапустите сервер (пересоздаст аудит-триггеры).')
        else:
            db.rollback()
            print('\nDRY-RUN: изменения НЕ записаны. Запустите с --apply для применения.')
    finally:
        db.close()


if __name__ == '__main__':
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    if not args:
        sys.exit('Укажите путь к Excel-файлу перечня оборудования.')
    run(args[0], apply='--apply' in sys.argv)

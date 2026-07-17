"""Тест импортёра Weld Balance на синтетическом файле.

Проверяет: парсинг листа Welds, нормализацию деталей в weld_point_part, dedup материалов,
best-effort связь с gun (по MNTC) и spot (по model+номер точки).
"""
import sqlite3

import openpyxl

from app.migrations import run_migrations
from scripts.import_weld_balance import run


def _make_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Welds'
    # Заголовок в строке 2 (валидатор проверяет col7/9/13/49).
    header = {7: '№ ПИСТОЛЕТА \n(GUN No.)', 9: '№ ТОЧКИ\n(SPOT №)',
              13: 'МАТЕРИАЛ 1\n(MATERIAL)', 49: 'X'}
    for col, val in header.items():
        ws.cell(row=2, column=col + 1, value=val)

    def put(r, mapping):
        for col, val in mapping.items():
            ws.cell(row=r, column=col + 1, value=val)

    # Строка 3: связывается с gun G.043 и spot 17; 2 детали.
    put(3, {6: 'PSW', 7: 'UCH-X', 8: 'G.043', 9: 17, 10: 'LH',
            11: 'Part A', 12: '#A1', 13: 'Steel 1', 15: 1.5,
            16: 'Part B', 17: '#B1', 18: 'Steel 2', 20: 0.8,
            49: 100, 50: 200, 51: 300})
    # Строка 4: gun/spot не совпадут; 1 деталь, материал 'Steel 1' (dedup).
    put(4, {6: 'CO2', 7: 'UCH-Y', 8: 'G.999', 9: 999, 10: 'RH',
            11: 'Part C', 12: '#C1', 13: 'Steel 1', 15: 2.0})
    wb.save(path)


def test_import_weld_balance(tmp_path):
    db_path = tmp_path / 'wb.db'
    run_migrations(str(db_path))
    con = sqlite3.connect(db_path)
    # минимальные справочники: модель A13T, пистолет g_num=43, точка №17
    con.execute("INSERT INTO brand (UniqueID, brand) VALUES (1,'Chery')")
    con.execute("INSERT INTO model (UniqueID, model_name, model_code, type, brand_id) "
                "VALUES (1,'Tiggo2','A13T','single',1)")
    con.execute("INSERT INTO gun (UniqueID, g_num, gun_type) VALUES (43,43,'UCH-X')")
    con.execute("INSERT INTO spot (UniqueID, spot_number, model_id) VALUES (14,17,1)")
    con.commit()
    con.close()

    wb_dir = tmp_path / 'wb'
    wb_dir.mkdir()
    _make_workbook(wb_dir / 'Weld_Balance_Table_A13T test.xlsx')

    run(str(wb_dir), str(db_path), apply=True)

    con = sqlite3.connect(db_path)
    con.row_factory = sqlite3.Row
    assert con.execute("SELECT COUNT(*) FROM weld_point").fetchone()[0] == 2
    assert con.execute("SELECT COUNT(*) FROM weld_point_part").fetchone()[0] == 3  # 2 + 1
    assert con.execute("SELECT COUNT(*) FROM wb_material").fetchone()[0] == 2      # Steel 1/2 (dedup)

    linked = con.execute("SELECT * FROM weld_point WHERE spot_number='17'").fetchone()
    assert linked['model_code'] == 'A13T'
    assert linked['gun_id'] == 43        # по MNTC G.043
    assert linked['spot_id'] == 14       # по (модель кода, spot_number)
    assert linked['coord_x'] == '100' and linked['coord_z'] == '300'

    parts = con.execute("SELECT layer_no, part_name, thickness FROM weld_point_part "
                        "WHERE weld_point_id=? ORDER BY layer_no", (linked['id'],)).fetchall()
    assert [p['part_name'] for p in parts] == ['Part A', 'Part B']

    # несвязанная точка: gun/spot NULL, но строка сохранена
    unl = con.execute("SELECT gun_id, spot_id FROM weld_point WHERE spot_number='999'").fetchone()
    assert unl['gun_id'] is None and unl['spot_id'] is None
    con.close()
